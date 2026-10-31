"""Финансовая модель «Мандата» на формулах Excel: эффект для ВТБ, рынок, три сценария, чувствительность, проверки.

    python submission/analytics/build_model.py             # собрать, пересчитать в LibreOffice, проверить
    python submission/analytics/build_model.py --no-recalc # только записать формулы

Параметры эффекта для ВТБ берутся из src/mandat/data/effect_scenarios.json (тот же файл читают калькулятор
и презентация), остальные допущения — из data/inputs.json. Результат: out/Mandat_financial_model.xlsx и
out/model_values.json — значения после пересчёта, из которых build_docs.py берёт числа для документов.

Устройство книги — по принципам FAST: допущения отдельно от расчётов, в строке одна формула для всех лет,
константы в столбце D, годы в столбцах F–J, 2026 год (столбец E) — исходные остатки. Три сценария
считаются на трёх листах одинаковой структуры, поэтому сравнение сценариев живое. «Торнадо» считается
блоками той же структуры на листе «Чувствительность».
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Callable
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]


def registry_values() -> dict[str, list[int]]:
    """Эталон листа «Проверки»: ставки и ущерб из строк C-15 и C-16 реестра утверждений."""
    text = (ROOT / "CLAIMS_REGISTRY.md").read_text(encoding="utf-8")
    out: dict[str, list[int]] = {}
    for key, row in (("fte_round", "C-15"), ("prevented_round", "C-16")):
        line = next(ln for ln in text.splitlines() if ln.startswith(f"| {row} |"))
        found = re.search(r"(\d+) / (\d+) / (\d+)", line)
        if found is None:
            raise SystemExit(f"в строке {row} реестра утверждений нет трёх сценариев «a / b / c»")
        out[key] = [int(x) for x in found.groups()]
    return out


OUT = HERE / "out"
XLSX = OUT / "Mandat_financial_model.xlsx"
VALUES = OUT / "model_values.json"

FONT = "Arial"
NAVY, COBALT, TEXT, MUTED, BG, LINE = "0A2973", "004CBC", "0B1F44", "5E6D8C", "F3F6FB", "DCE4F0"
INPUT, LINK, KEY = "0000FF", "008000", "FFFF00"
OK_FILL, BAD_FILL, WARN_FILL = "E2F4EB", "FDE6E7", "FFF1D4"

MLN = '#,##0.0;(#,##0.0);"–"'
MLN2 = '#,##0.00;(#,##0.00);"–"'
MLN3 = '#,##0.000;(#,##0.000);"–"'
INT = '#,##0;(#,##0);"–"'
NUM1 = '#,##0.0;(#,##0.0);"–"'
PCT = '0.0%;(0.0%);"–"'
PCT0 = '0%;(0%);"–"'
YEAR = "0"
MONTHS = '0.0;(0.0);"–"'

YEAR_COLS = ["F", "G", "H", "I", "J"]
START_COL, CONST_COL = "E", "D"
SCEN_COLS = ["D", "E", "F"]
SHEET_CALC = ["Консервативный", "Базовый", "Оптимистичный"]


def q(sheet: str) -> str:
    """Имя листа для ссылки: в кавычках, если в нём пробел."""
    return f"'{sheet}'" if " " in sheet else sheet


def prev(col: str) -> str:
    return chr(ord(col) - 1)


def font(color: str = "000000", bold: bool = False, size: float = 10, italic: bool = False) -> Font:
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


THIN = Side(style="thin", color=LINE)
BOTTOM = Border(bottom=THIN)


def decimals(values: list[float]) -> int:
    def one(v: float) -> int:
        s = repr(float(v))
        return 0 if s.endswith(".0") else min(3, len(s.split(".")[1]))

    return max(one(v) for v in values)


def fmt_for(unit: str, values: list[float]) -> str:
    d = decimals(values)
    if unit.startswith("доля"):
        return PCT if d >= 3 else PCT0
    if unit == "год":
        return YEAR
    if d == 0:
        return INT
    return {1: MLN, 2: MLN2, 3: MLN3}[d]


# ---------------------------------------------------------------- данные


def load() -> tuple[dict, dict]:
    data = json.loads((HERE / "data" / "inputs.json").read_text(encoding="utf-8"))
    effect = json.loads((ROOT / "src" / "mandat" / "data" / "effect_scenarios.json").read_text(encoding="utf-8"))
    names = effect["names"]
    if names != data["scenarios"]:
        raise SystemExit(f"названия сценариев расходятся: {names} и {data['scenarios']}")
    three = lambda v: v if isinstance(v, list) else [v, v, v]  # noqa: E731
    presentation = {
        "code": "П",
        "title": "Эффект для ВТБ: параметры презентации (effect_scenarios.json)",
        "items": [
            {
                "id": "clients",
                "name": "Клиенты ВТБ — предприятия и предприниматели",
                "unit": "млн",
                "values": [effect["clients"] / 1e6] * 3,
                "type": "факт",
                "key": True,
                "source": "Пресс-релиз Группы ВТБ к МСФО за I полугодие 2026, 29.07.2026 (M-09, S05)",
                "check": "внутренняя цифра банка",
            },
            {
                "id": "event_share",
                "name": "Доля клиентов со сменой или добавлением подписанта в год",
                "unit": "доля",
                "values": three(effect["eventShare"]),
                "type": "допущение",
                "key": True,
                "source": "Допущение команды (C-15); публичных данных нет",
                "check": "число событий за год в журнале банка",
            },
            {
                "id": "minutes",
                "name": "Минут на кейс вручную, вместе с разбором",
                "unit": "минут",
                "values": three(effect["minutesSaved"]),
                "type": "допущение",
                "key": True,
                "source": "Допущение команды (C-15); замер — в пилоте",
                "check": "хронометраж в пилоте",
            },
            {
                "id": "hours_fte",
                "name": "Рабочих часов на ставку в год",
                "unit": "часов",
                "values": three(effect["hoursPerFte"]),
                "type": "допущение",
                "source": "Около 1 970 часов по производственному календарю минус отпуск",
                "check": "норма банка",
            },
            {
                "id": "incidents_q",
                "name": "Операций без согласия у юрлиц в среднем за квартал I полугодия 2026",
                "unit": "операций",
                "values": three(effect["incidentsPerQuarter"]),
                "type": "факт",
                "source": "Банк России: I кв. — 818, II кв. — 780 операций (M-07, M-15, S03, S04)",
                "check": "—",
            },
            {
                "id": "loss_q",
                "name": "Объём таких операций в среднем за квартал I полугодия 2026",
                "unit": "млн ₽",
                "values": three(effect["lossPerQuarterMln"]),
                "type": "факт",
                "source": "Банк России: I кв. — 573, II кв. — 831,4 млн ₽ (M-07, M-15, S03, S04)",
                "check": "—",
            },
            {
                "id": "vtb_share",
                "name": "Доля ВТБ в таких инцидентах",
                "unit": "доля",
                "values": three(effect["vtbShare"]),
                "type": "допущение",
                "source": "Допущение команды (C-16)",
                "check": "данные антифрода банка",
            },
            {
                "id": "authority_share",
                "name": "Доля инцидентов, связанных с полномочиями",
                "unit": "доля",
                "values": three(effect["authorityShare"]),
                "type": "допущение",
                "source": "Допущение команды (C-16)",
                "check": "разметка инцидентов банка",
            },
        ],
    }
    data["sections"] = [presentation, *data["sections"]]
    return data, effect


# ---------------------------------------------------------------- оформление


def title(ws: Worksheet, text: str, sub: str) -> None:
    ws["A1"] = text
    ws["A1"].font = font(NAVY, True, 14)
    ws["A2"] = sub
    ws["A2"].font = font(MUTED, size=9, italic=True)
    ws.sheet_view.showGridLines = False


def header(ws: Worksheet, row: int, labels: dict[str, str]) -> None:
    for col, text in labels.items():
        c = ws[f"{col}{row}"]
        c.value = text
        c.font = font("FFFFFF", True)
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center" if col not in "AB" else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def section(ws: Worksheet, row: int, text: str, last: str) -> None:
    ws[f"A{row}"] = text
    for col in range(1, ord(last) - 63):
        cell = ws.cell(row, col)
        cell.fill = fill(BG)
        cell.font = font(NAVY, True)


def widths(ws: Worksheet, spec: dict[str, float]) -> None:
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def page(ws: Worksheet) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def style_value(cell, kind: str, fmt: str) -> None:
    """kind: input — введено руками, link — ссылка на другой лист, calc — формула."""
    color = {"input": INPUT, "link": LINK, "calc": "000000"}[kind]
    cell.font = font(color)
    cell.number_format = fmt
    cell.alignment = Alignment(horizontal="right")


# ---------------------------------------------------------------- допущения


class Inputs:
    """Лист «Допущения»: строка на параметр, столбцы D–F — сценарии."""

    def __init__(self, wb: Workbook, data: dict) -> None:
        self.ws = ws = wb.create_sheet("Допущения")
        self.rows: dict[str, int] = {}
        self.codes: dict[str, str] = {}
        self.items: dict[str, dict] = {}
        title(
            ws,
            "Допущения",
            "Синий — введено руками, жёлтая заливка — ключевые допущения, которые банк заменит своими данными",
        )
        ws["A3"] = "Год начала модели (исходные остатки)"
        ws["A3"].font = font(TEXT, True)
        ws["D3"] = data["start_year"]
        style_value(ws["D3"], "input", YEAR)
        self.rows["start_year"] = 3
        header(
            ws,
            5,
            {
                "A": "№",
                "B": "Параметр",
                "C": "Единица",
                "D": data["scenarios"][0],
                "E": data["scenarios"][1],
                "F": data["scenarios"][2],
                "G": "Тип",
                "H": "Источник или обоснование",
                "I": "Как проверить на пилоте",
            },
        )
        r = 6
        for sec in data["sections"]:
            section(ws, r, f"{sec['code']}. {sec['title']}", "I")
            r += 1
            for i, item in enumerate(sec["items"], start=1):
                code = f"{sec['code']}{i}"
                self.rows[item["id"]], self.codes[item["id"]], self.items[item["id"]] = r, code, item
                ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = code, item["name"], item["unit"]
                f = fmt_for(item["unit"], item["values"])
                for col, v in zip(SCEN_COLS, item["values"], strict=True):
                    ws[f"{col}{r}"] = v
                    style_value(ws[f"{col}{r}"], "input", f)
                    if item.get("key"):
                        ws[f"{col}{r}"].fill = fill(KEY)
                ws[f"G{r}"], ws[f"H{r}"], ws[f"I{r}"] = item["type"], item["source"], item["check"]
                for col in "ABCGHI":
                    ws[f"{col}{r}"].font = font(TEXT if col in "AB" else MUTED, size=10 if col in "AB" else 9)
                    ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
                for col in "ABCDEFGHI":
                    ws[f"{col}{r}"].border = BOTTOM
                r += 1
        widths(ws, {"A": 6, "B": 46, "C": 18, "D": 15, "E": 15, "F": 15, "G": 12, "H": 70, "I": 34})
        ws.freeze_panes = "D6"
        page(ws)

    def ref(self, key: str, scen: int) -> str:
        return f"Допущения!${SCEN_COLS[scen]}${self.rows[key]}"


class Plan:
    """Лист «План»: договоры, численность и оклады по годам для каждого сценария."""

    def __init__(self, wb: Workbook, data: dict) -> None:
        self.ws = ws = wb.create_sheet("План")
        self.rows: dict[tuple[int, str], int] = {}
        title(
            ws,
            "План продаж и команды",
            "Новые договоры и численность в ставках на конец года; 2026 — исходные значения",
        )
        header(
            ws,
            4,
            {
                "A": "№",
                "B": "Показатель",
                "C": "Единица",
                "D": "Оклад 2027",
                "E": "2026",
                **{c: str(y) for c, y in zip(YEAR_COLS, data["years"], strict=True)},
                "K": "Итого",
            },
        )
        r = 5
        plan = data["plan"]
        for s, scen in enumerate(data["scenarios"]):
            section(ws, r, scen, "K")
            r += 1
            for j, deal in enumerate(plan["deals"], start=1):
                self._row(r, f"{scen[0]}Д{j}", deal["name"], deal["unit"], None, 0, deal["values"][s], INT)
                ws[f"K{r}"] = f"=SUM(F{r}:J{r})"
                style_value(ws[f"K{r}"], "calc", INT)
                self.rows[(s, deal["id"])] = r
                r += 1
            for j, role in enumerate(plan["roles"], start=1):
                self._row(
                    r,
                    f"{scen[0]}Ч{j}",
                    role["name"],
                    "ставок",
                    role["salary"][s],
                    role["start"],
                    role["values"][s],
                    NUM1,
                )
                self.rows[(s, role["id"])] = r
                r += 1
            self.rows[(s, "total_first")] = self.rows[(s, plan["roles"][0]["id"])]
            r += 1
        ws[f"A{r}"] = plan["salary_source"]
        ws[f"A{r}"].font = font(MUTED, size=9, italic=True)
        widths(ws, {"A": 7, "B": 50, "C": 10, "D": 12, "E": 9, **dict.fromkeys(YEAR_COLS, 9), "K": 9})
        ws.freeze_panes = "F5"
        page(ws)

    def _row(
        self, r: int, code: str, name: str, unit: str, salary: float | None, start: float, values: list[float], f: str
    ) -> None:
        ws = self.ws
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = code, name, unit
        for col in "ABC":
            ws[f"{col}{r}"].font = font(TEXT if col != "C" else MUTED)
        if salary is not None:
            ws[f"D{r}"] = salary
            style_value(ws[f"D{r}"], "input", INT)
        ws[f"E{r}"] = start
        style_value(ws[f"E{r}"], "input", f)
        for col, v in zip(YEAR_COLS, values, strict=True):
            ws[f"{col}{r}"] = v
            style_value(ws[f"{col}{r}"], "input", f)
        for col in "ABCDEFGHIJK":
            ws[f"{col}{r}"].border = BOTTOM

    def ref(self, scen: int, key: str, col: str) -> str:
        return f"План!{col}{self.rows[(scen, key)]}"

    def salary(self, scen: int, key: str) -> str:
        return f"План!$D${self.rows[(scen, key)]}"


# ---------------------------------------------------------------- лист с тремя сценариями в столбцах


class Columns:
    """Лист, где сценарии стоят в столбцах D–F (эффект для банка, рынок)."""

    def __init__(self, wb: Workbook, name: str, heading: str, sub: str, scenarios: list[str]) -> None:
        self.ws = ws = wb.create_sheet(name)
        self.name = name
        self.rows: dict[str, int] = {}
        title(ws, heading, sub)
        header(
            ws,
            4,
            {
                "A": "№",
                "B": "Показатель",
                "C": "Единица",
                "D": scenarios[0],
                "E": scenarios[1],
                "F": scenarios[2],
                "G": "Как считается",
            },
        )
        self.r = 5
        self.n = 0
        widths(ws, {"A": 6, "B": 54, "C": 16, "D": 15, "E": 15, "F": 15, "G": 80})
        ws.freeze_panes = "D5"
        page(ws)

    def section(self, text: str) -> None:
        section(self.ws, self.r, text, "G")
        self.r += 1

    def row(
        self,
        key: str,
        label: str,
        unit: str,
        formula: Callable[[int, str], str],
        fmt: str,
        note: str = "",
        kind: str = "calc",
    ) -> None:
        ws, r = self.ws, self.r
        self.n += 1
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"G{r}"] = self.n, label, unit, note
        ws[f"A{r}"].font, ws[f"B{r}"].font, ws[f"C{r}"].font = font(MUTED), font(TEXT), font(MUTED)
        ws[f"G{r}"].font = font(MUTED, size=9)
        ws[f"G{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        for s, col in enumerate(SCEN_COLS):
            ws[f"{col}{r}"] = formula(s, col)
            style_value(ws[f"{col}{r}"], kind, fmt)
        for col in "ABCDEFG":
            ws[f"{col}{r}"].border = BOTTOM
        self.rows[key] = r
        self.r += 1

    def at(self, key: str, col: str) -> str:
        return f"{col}{self.rows[key]}"

    def ext(self, key: str, scen: int) -> str:
        return f"{q(self.name)}!${SCEN_COLS[scen]}${self.rows[key]}"


def bank_sheet(wb: Workbook, inp: Inputs, data: dict) -> Columns:
    b = Columns(
        wb,
        "Эффект для банка",
        "Эффект и окупаемость для ВТБ",
        "Ставки и ущерб повторяют расчёт презентации (C-15, C-16); деньги — млн ₽ в год без НДС, цены 2027 года",
        data["scenarios"],
    )
    L = lambda key: lambda s, col: f"={inp.ref(key, s)}"  # noqa: E731
    a = b.at

    b.section("Параметры (ссылки на «Допущения»)")
    for key in [
        "clients",
        "event_share",
        "minutes",
        "hours_fte",
        "incidents_q",
        "loss_q",
        "vtb_share",
        "authority_share",
        "salary_ops",
        "bank_contrib",
        "overhead",
        "realization",
        "bank_loss_share",
        "ramp",
        "fixed_t1",
        "price_client",
        "impl_t1",
        "vat",
        "bank_it_cost",
        "bank_integration_py",
        "bank_infra",
        "bank_security",
        "bank_support_py",
        "build_py",
        "build_team",
        "build_horizon",
    ]:
        item = inp.items[key]
        b.row(
            key,
            item["name"],
            item["unit"],
            L(key),
            fmt_for(item["unit"], item["values"]),
            f"Допущения, строка {inp.codes[key]}",
            kind="link",
        )

    b.section("1. Высвобождение времени")
    b.row(
        "events",
        "Кейсов смены или добавления подписанта в год",
        "тыс.",
        lambda s, c: f"={a('clients', c)}*1000*{a('event_share', c)}",
        MLN,
        "клиенты × доля с событием",
    )
    b.row(
        "hours",
        "Часов ручной проверки в год",
        "тыс. часов",
        lambda s, c: f"={a('events', c)}*{a('minutes', c)}/60",
        MLN,
        "кейсы × минуты / 60",
    )
    b.row(
        "fte",
        "Высвобождается ставок",
        "ставок",
        lambda s, c: f"={a('hours', c)}*1000/{a('hours_fte', c)}",
        NUM1,
        "часы / часы на ставку",
    )
    b.row(
        "fte_round",
        "То же, округлено (как в презентации, C-15)",
        "ставок",
        lambda s, c: f"=ROUND({a('fte', c)},0)",
        INT,
        "сверяется на листе «Проверки»",
    )
    b.row(
        "cost_fte",
        "Полная стоимость ставки",
        "млн ₽ в год",
        lambda s, c: f"={a('salary_ops', c)}*12*(1+{a('bank_contrib', c)})*(1+{a('overhead', c)})/1000",
        MLN2,
        "оклад × 12 × (1 + взносы) × (1 + накладные)",
    )
    b.row(
        "capacity",
        "Стоимость высвобожденного времени (ёмкость)",
        "млн ₽ в год",
        lambda s, c: f"={a('fte', c)}*{a('cost_fte', c)}",
        MLN,
        "ставки × стоимость ставки",
    )
    b.row(
        "cash_saving",
        "Денежная экономия на времени",
        "млн ₽ в год",
        lambda s, c: f"={a('capacity', c)}*{a('realization', c)}",
        MLN,
        "ёмкость × доля, которая становится экономией",
    )

    b.section("2. Предотвращённый ущерб")
    b.row(
        "loss_op",
        "Средняя операция без согласия у юрлица",
        "млн ₽",
        lambda s, c: f"={a('loss_q', c)}/{a('incidents_q', c)}",
        MLN2,
        "объём за квартал / число операций",
    )
    b.row(
        "ops_year",
        "Таких операций в год по всем банкам",
        "операций",
        lambda s, c: f"={a('incidents_q', c)}*4",
        INT,
        "квартал × 4",
    )
    b.row(
        "prevented",
        "Предотвращённый ущерб клиентов ВТБ",
        "млн ₽ в год",
        lambda s, c: f"={a('ops_year', c)}*{a('vtb_share', c)}*{a('authority_share', c)}*{a('loss_op', c)}",
        MLN,
        "операции × доля ВТБ × доля с полномочиями × средняя операция",
    )
    b.row(
        "prevented_round",
        "То же, округлено (как в презентации, C-16)",
        "млн ₽ в год",
        lambda s, c: f"=ROUND({a('prevented', c)},0)",
        INT,
        "сверяется на листе «Проверки»",
    )
    b.row(
        "bank_loss",
        "Из него несёт банк",
        "млн ₽ в год",
        lambda s, c: f"={a('prevented', c)}*{a('bank_loss_share', c)}",
        MLN,
        "ущерб × доля банка",
    )

    b.section("3. Эффект в год")
    b.row(
        "effect_capacity",
        "Прямой эффект банка: ёмкость",
        "млн ₽ в год",
        lambda s, c: f"={a('capacity', c)}+{a('bank_loss', c)}",
        MLN,
        "время целиком + ущерб, который несёт банк",
    )
    b.row(
        "effect_cash",
        "Прямой эффект банка: деньги",
        "млн ₽ в год",
        lambda s, c: f"={a('cash_saving', c)}+{a('bank_loss', c)}",
        MLN,
        "денежная экономия + ущерб банка",
    )
    b.row(
        "effect_social",
        "Эффект с ущербом клиентов",
        "млн ₽ в год",
        lambda s, c: f"={a('capacity', c)}+{a('prevented', c)}",
        MLN,
        "время + весь предотвращённый ущерб",
    )

    b.section("4. Затраты банка")
    b.row(
        "sub_vtb",
        "Подписка «Мандата» для ВТБ",
        "млн ₽ в год",
        lambda s, c: f"={a('fixed_t1', c)}+{a('clients', c)}*{a('price_client', c)}",
        MLN,
        "постоянная часть + клиенты × плата за клиента; лицензия из реестра ПО — без НДС",
    )
    b.row(
        "impl_vat",
        "Внедрение с НДС",
        "млн ₽",
        lambda s, c: f"={a('impl_t1', c)}*(1+{a('vat', c)})",
        MLN,
        "услуга облагается НДС, банк его не вычитает",
    )
    b.row(
        "integration",
        "Интеграция силами банка",
        "млн ₽",
        lambda s, c: f"={a('bank_integration_py', c)}*{a('bank_it_cost', c)}",
        MLN,
        "человеко-годы × стоимость",
    )
    b.row(
        "one_off",
        "Разовые затраты банка",
        "млн ₽",
        lambda s, c: f"={a('impl_vat', c)}+{a('integration', c)}+{a('bank_infra', c)}+{a('bank_security', c)}",
        MLN,
        "внедрение + интеграция + инфраструктура + оценка ИБ",
    )
    b.row(
        "annual_cost",
        "Ежегодные затраты банка",
        "млн ₽ в год",
        lambda s, c: f"={a('sub_vtb', c)}+{a('bank_support_py', c)}*{a('bank_it_cost', c)}",
        MLN,
        "подписка + сопровождение силами банка",
    )

    for view, label in (("cash", "деньги"), ("capacity", "ёмкость")):
        eff = f"effect_{view}"
        b.section(f"5. Окупаемость для банка: {label}")
        b.row(
            f"y1_{view}",
            "Чистый эффект первого года",
            "млн ₽",
            lambda s, c, e=eff: f"={a(e, c)}*{a('ramp', c)}-{a('annual_cost', c)}",
            MLN,
            "эффект × раскачка − ежегодные",
        )
        b.row(
            f"y2_{view}",
            "Чистый эффект со второго года",
            "млн ₽ в год",
            lambda s, c, e=eff: f"={a(e, c)}-{a('annual_cost', c)}",
            MLN,
            "эффект − ежегодные",
        )
        b.row(
            f"roi_{view}",
            "Отдача за 3 года на вложенный рубль",
            "раз",
            lambda s, c, v=view: (
                f"=({a(f'y1_{v}', c)}+2*{a(f'y2_{v}', c)}-{a('one_off', c)})"
                f"/({a('one_off', c)}+3*{a('annual_cost', c)})"
            ),
            '0.0"×";(0.0"×");"–"',
            "(чистый эффект 3 лет − разовые) / (разовые + ежегодные × 3)",
        )
        b.row(
            f"payback_{view}",
            "Срок окупаемости разовых затрат",
            "месяцев",
            lambda s, c, v=view: (
                f"=IF({a(f'y1_{v}', c)}>={a('one_off', c)},12*{a('one_off', c)}/{a(f'y1_{v}', c)},"
                f'IF({a(f"y2_{v}", c)}<=0,"не окупается",'
                f"12+12*({a('one_off', c)}-{a(f'y1_{v}', c)})/{a(f'y2_{v}', c)}))"
            ),
            MONTHS,
            "месяцы до возврата разовых затрат с учётом раскачки",
        )

    b.section("6. Цена и «купить или сделать»")
    b.row(
        "price_share",
        "Подписка в доле прямого эффекта (ёмкость)",
        "доля",
        lambda s, c: f"={a('sub_vtb', c)}/{a('effect_capacity', c)}",
        PCT,
        "условие цены: не больше 15 %",
    )
    b.row(
        "price_share_cash",
        "Подписка в доле прямого эффекта (деньги)",
        "доля",
        lambda s, c: f"={a('sub_vtb', c)}/{a('effect_cash', c)}",
        PCT,
        "справочно",
    )
    b.row(
        "buy_annual",
        "Купить: затраты в год за горизонт",
        "млн ₽ в год",
        lambda s, c: f"={a('annual_cost', c)}+{a('one_off', c)}/{a('build_horizon', c)}",
        MLN,
        "ежегодные + разовые / горизонт",
    )
    b.row(
        "build_annual",
        "Сделать самим: затраты в год за горизонт",
        "млн ₽ в год",
        lambda s, c: (
            f"={a('build_py', c)}*{a('bank_it_cost', c)}/{a('build_horizon', c)}"
            f"+{a('build_team', c)}*{a('bank_it_cost', c)}+{a('bank_infra', c)}/{a('build_horizon', c)}"
        ),
        MLN,
        "создание / горизонт + команда сопровождения + инфраструктура / горизонт",
    )
    b.row(
        "buy_vs_build",
        "Купить в доле от «сделать самим»",
        "доля",
        lambda s, c: f"={a('buy_annual', c)}/{a('build_annual', c)}",
        PCT,
        "условие цены: меньше 100 %",
    )
    return b


def market_sheet(wb: Workbook, inp: Inputs, bank: Columns, data: dict) -> Columns:
    m = Columns(
        wb,
        "Рынок",
        "Объём рынка: банки России",
        "Снизу вверх — банки по уровням × цена подписки; сверху вниз — ценность задачи × доля поставщика. "
        "Млн ₽ в год без НДС, цены 2027 года",
        data["scenarios"],
    )
    L = lambda key: lambda s, col: f"={inp.ref(key, s)}"  # noqa: E731
    a = m.at

    m.section("Параметры (ссылки на «Допущения» и «Эффект для банка»)")
    for key in [
        "banks_total",
        "banks_t1",
        "banks_t2",
        "vtb_rel_share",
        "t1_rel_share",
        "t2_rel_share",
        "price_t3",
        "capture",
        "clients",
        "fixed_t1",
        "fixed_t2",
        "price_client",
        "expansion",
        "incidents_q",
        "authority_share",
        "bank_loss_share",
    ]:
        item = inp.items[key]
        m.row(
            key,
            item["name"],
            item["unit"],
            L(key),
            fmt_for(item["unit"], item["values"]),
            f"Допущения, строка {inp.codes[key]}",
            kind="link",
        )
    for key, label in [
        ("capacity", "Стоимость высвобожденного времени ВТБ"),
        ("loss_op", "Средняя операция без согласия"),
    ]:
        m.row(
            key,
            label,
            "млн ₽ в год" if key == "capacity" else "млн ₽",
            lambda s, c, k=key: f"={bank.ext(k, s)}",
            MLN if key == "capacity" else MLN2,
            "лист «Эффект для банка»",
            kind="link",
        )

    m.section("1. Банки и клиенты-бизнесы")
    m.row(
        "banks_t3",
        "Уровень 3: остальные банки",
        "банков",
        lambda s, c: f"={a('banks_total', c)}-{a('banks_t1', c)}-{a('banks_t2', c)}",
        INT,
        "всего − уровень 1 − уровень 2",
    )
    m.row(
        "rel_total",
        "Отношений «банк — клиент-бизнес» в России",
        "млн",
        lambda s, c: f"={a('clients', c)}/{a('vtb_rel_share', c)}",
        MLN2,
        "клиенты ВТБ / доля ВТБ",
    )
    m.row(
        "rel_t1",
        "Из них в банках уровня 1",
        "млн",
        lambda s, c: f"={a('rel_total', c)}*{a('t1_rel_share', c)}",
        MLN2,
        "всего × доля уровня 1",
    )
    m.row(
        "rel_t2",
        "В банках уровня 2",
        "млн",
        lambda s, c: f"={a('rel_total', c)}*{a('t2_rel_share', c)}",
        MLN2,
        "всего × доля уровня 2",
    )
    m.row(
        "rel_t3",
        "В банках уровня 3",
        "млн",
        lambda s, c: f"={a('rel_total', c)}-{a('rel_t1', c)}-{a('rel_t2', c)}",
        MLN2,
        "остаток",
    )
    m.row(
        "avg_t1",
        "Клиентов-бизнесов у банка уровня 1, кроме ВТБ, в среднем",
        "млн",
        lambda s, c: f"=({a('rel_t1', c)}-{a('clients', c)})/({a('banks_t1', c)}-1)",
        MLN3,
        "(уровень 1 − ВТБ) / (банков уровня 1 − 1); база для подписки других банков",
    )
    m.row(
        "avg_t2",
        "Клиентов-бизнесов у банка уровня 2 в среднем",
        "млн",
        lambda s, c: f"={a('rel_t2', c)}/{a('banks_t2', c)}",
        MLN3,
        "уровень 2 / банков уровня 2",
    )

    m.section("2. Снизу вверх: все банки по прайсу")
    m.row(
        "tam_t1",
        "Уровень 1",
        "млн ₽ в год",
        lambda s, c: f"={a('banks_t1', c)}*{a('fixed_t1', c)}+{a('rel_t1', c)}*{a('price_client', c)}",
        MLN,
        "банки × постоянная часть + клиенты × плата за клиента",
    )
    m.row(
        "tam_t2",
        "Уровень 2",
        "млн ₽ в год",
        lambda s, c: f"={a('banks_t2', c)}*{a('fixed_t2', c)}+{a('rel_t2', c)}*{a('price_client', c)}",
        MLN,
        "то же",
    )
    m.row(
        "tam_t3",
        "Уровень 3 (через партнёров)",
        "млн ₽ в год",
        lambda s, c: f"={a('banks_t3', c)}*{a('price_t3', c)}",
        MLN,
        "банки × цена через партнёра",
    )
    m.row(
        "tam",
        "TAM: все банки России",
        "млн ₽ в год",
        lambda s, c: f"={a('tam_t1', c)}+{a('tam_t2', c)}+{a('tam_t3', c)}",
        MLN,
        "сумма уровней",
    )
    m.row(
        "sam",
        "SAM: банки уровней 1 и 2 (установка в контуре)",
        "млн ₽ в год",
        lambda s, c: f"={a('tam_t1', c)}+{a('tam_t2', c)}",
        MLN,
        "уровни, которые обслуживаем сами",
    )

    m.row(
        "full_mult",
        "Подписка при полном охвате операций, к одной операции",
        "раз",
        lambda s, c: f"=(1+{a('expansion', c)})^4",
        '0.00"×"',
        "банк добавляет операции четыре года: (1 + рост подписки) ^ 4",
    )
    m.row(
        "tam_full",
        "TAM при полном охвате операций",
        "млн ₽ в год",
        lambda s, c: f"={a('tam', c)}*{a('full_mult', c)}",
        MLN,
        "TAM × множитель",
    )
    m.row(
        "sam_full",
        "SAM при полном охвате операций",
        "млн ₽ в год",
        lambda s, c: f"={a('sam', c)}*{a('full_mult', c)}",
        MLN,
        "SAM × множитель",
    )

    m.section("3. SOM: годовая подписка на конец 2031 года по плану продаж")
    m.row(
        "som_nominal",
        "Подписка на конец 2031 года (ARR), номинал",
        "млн ₽ в год",
        lambda s, c: f"={SHEET_CALC[s]}!$J${{arr}}",
        MLN,
        "лист сценария, строка «Годовая подписка на конец года»",
        kind="link",
    )
    m.row(
        "index_2031",
        "Индекс цен 2031 к 2027",
        "раз",
        lambda s, c: f"={SHEET_CALC[s]}!$J${{idx}}",
        "0.000",
        "лист сценария",
        kind="link",
    )
    m.row(
        "som",
        "SOM в ценах 2027 года",
        "млн ₽ в год",
        lambda s, c: f"={a('som_nominal', c)}/{a('index_2031', c)}",
        MLN,
        "номинал / индекс",
    )
    m.row(
        "som_share",
        "SOM в доле SAM при полном охвате",
        "доля",
        lambda s, c: f"={a('som', c)}/{a('sam_full', c)}",
        PCT,
        "SOM включает рост подписки за счёт новых операций, поэтому сравнивается с SAM при полном охвате",
    )

    m.section("4. Сверху вниз: ценность задачи для всех банков")
    m.row(
        "pool_time",
        "Стоимость ручной проверки во всех банках",
        "млн ₽ в год",
        lambda s, c: f"={a('capacity', c)}/{a('vtb_rel_share', c)}",
        MLN,
        "ёмкость ВТБ / доля ВТБ в отношениях",
    )
    m.row(
        "pool_loss",
        "Ущерб клиентов из-за полномочий во всех банках",
        "млн ₽ в год",
        lambda s, c: f"={a('incidents_q', c)}*4*{a('authority_share', c)}*{a('loss_op', c)}",
        MLN,
        "операции в год × доля с полномочиями × средняя операция",
    )
    m.row(
        "pool_bank",
        "Ценность для банков: время + ущерб, который несут банки",
        "млн ₽ в год",
        lambda s, c: f"={a('pool_time', c)}+{a('pool_loss', c)}*{a('bank_loss_share', c)}",
        MLN,
        "",
    )
    m.row(
        "pool_social",
        "Ценность с ущербом клиентов",
        "млн ₽ в год",
        lambda s, c: f"={a('pool_time', c)}+{a('pool_loss', c)}",
        MLN,
        "",
    )
    m.row(
        "topdown",
        "Выручка поставщика сверху вниз",
        "млн ₽ в год",
        lambda s, c: f"={a('pool_bank', c)}*{a('capture', c)}",
        MLN,
        "ценность для банков × доля поставщика",
    )
    m.row(
        "tam_vs_pool",
        "TAM по прайсу в доле ценности для банков",
        "доля",
        lambda s, c: f"={a('tam', c)}/{a('pool_bank', c)}",
        PCT,
        "выше доли поставщика: постоянная часть подписки оплачивает сопровождение правил, а не только время",
    )
    return m


# ---------------------------------------------------------------- расчётный блок сценария


class Block:
    """Полный расчёт одного сценария: входы (ссылки) → договоры → выручка → затраты → отчёт → деньги.

    src(key) — выражение для константы, plan(key, col) — выражение для значения года. Блоки чувствительности
    подставляют в них множители, формулы ниже одинаковы во всех блоках.
    """

    def __init__(
        self,
        ws: Worksheet,
        top: int,
        data: dict,
        src: Callable[[str], str],
        plan: Callable[[str, str], str],
        salary: Callable[[str], str],
        start_year: str,
        avg: Callable[[str], str],
        header_row: bool = True,
    ) -> None:
        self.ws, self.r, self.data = ws, top, data
        self.rows: dict[str, int] = {}
        self.n = 0
        if header_row:
            header(
                ws,
                self.r,
                {
                    "A": "№",
                    "B": "Показатель",
                    "C": "Единица",
                    "D": "Константа",
                    "E": "2026",
                    **{c: str(y) for c, y in zip(YEAR_COLS, data["years"], strict=True)},
                    "K": "Итого 2027–2031",
                    "L": "Как считается",
                },
            )
            self.r += 1
        self.build(src, plan, salary, start_year, avg)

    # --- запись строк

    def _base(self, key: str, label: str, unit: str, note: str) -> int:
        ws, r = self.ws, self.r
        self.n += 1
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"L{r}"] = self.n, label, unit, note
        ws[f"A{r}"].font, ws[f"B{r}"].font, ws[f"C{r}"].font = font(MUTED), font(TEXT), font(MUTED)
        ws[f"L{r}"].font = font(MUTED, size=9)
        for col in "ABCDEFGHIJKL":
            ws[f"{col}{r}"].border = BOTTOM
        self.rows[key] = r
        self.r += 1
        return r

    def const(self, key: str, label: str, unit: str, expr: str, fmt: str, note: str = "", kind: str = "link") -> None:
        r = self._base(key, label, unit, note)
        self.ws[f"D{r}"] = expr
        style_value(self.ws[f"D{r}"], kind, fmt)

    def series(
        self,
        key: str,
        label: str,
        unit: str,
        f: Callable[[str], str],
        fmt: str,
        note: str = "",
        start: str | None = None,
        total: bool = False,
        kind: str = "calc",
    ) -> None:
        r = self._base(key, label, unit, note)
        if start is not None:
            self.ws[f"E{r}"] = start
            style_value(self.ws[f"E{r}"], "link" if start.startswith("=") and "!" in start else "calc", fmt)
        for col in YEAR_COLS:
            self.ws[f"{col}{r}"] = f(col)
            style_value(self.ws[f"{col}{r}"], kind, fmt)
        if total:
            self.ws[f"K{r}"] = f"=SUM(F{r}:J{r})"
            style_value(self.ws[f"K{r}"], "calc", fmt)
            self.ws[f"K{r}"].font = font(TEXT, True)

    def section(self, text: str) -> None:
        section(self.ws, self.r, text, "L")
        self.r += 1

    def c(self, key: str) -> str:
        return f"$D${self.rows[key]}"

    def s(self, key: str, col: str) -> str:
        return f"{col}{self.rows[key]}"

    def y(self, col: str) -> str:
        return f"{col}${self.rows['year']}"

    # --- расчёт

    def build(
        self,
        src: Callable[[str], str],
        plan: Callable[[str, str], str],
        salary: Callable[[str], str],
        start_year: str,
        avg: Callable[[str], str],
    ) -> None:
        c, s, y = self.c, self.s, self.y
        roles = self.data["plan"]["roles"]

        self.section("Входные данные (ссылки на «Допущения», «План», «Рынок»)")
        consts = [
            ("vtb_year", "Год договора с ВТБ", "год", YEAR),
            ("clients", "Клиенты-бизнесы ВТБ", "млн", MLN2),
            ("fixed_t1", "Подписка уровня 1: постоянная часть", "млн ₽ в год", MLN),
            ("price_client", "Плата за активного клиента-бизнес", "₽ в год", MLN),
            ("fixed_t2", "Подписка уровня 2: постоянная часть", "млн ₽ в год", MLN),
            ("impl_t1", "Внедрение, уровень 1", "млн ₽", MLN),
            ("impl_t2", "Внедрение, уровень 2", "млн ₽", MLN),
            ("indexation", "Индексация цен", "доля в год", PCT),
            ("expansion", "Рост подписки за счёт новых операций", "доля в год", PCT),
            ("raise", "Рост окладов", "доля в год", PCT),
            ("contrib_low", "Взносы в пределах базы", "доля", PCT),
            ("contrib_high", "Взносы сверх базы", "доля", PCT),
            ("contrib_injury", "Взнос на травматизм", "доля", PCT),
            ("contrib_base", "Предельная база взносов 2026", "млн ₽ в год", MLN3),
            ("base_growth", "Рост базы", "доля в год", PCT),
            ("sec_first", "Анализ уязвимостей по ОУД4 — первый", "млн ₽", MLN),
            ("sec_annual", "Повторная оценка релизов", "млн ₽ в год", MLN),
            ("legal_ext", "Внешняя юридическая экспертиза правил", "млн ₽ в год", MLN),
            ("infra_head", "Инфраструктура на сотрудника", "млн ₽ в год", MLN2),
            ("sales_fixed", "Продажи: постоянная часть", "млн ₽ в год", MLN),
            ("sales_deal", "Продажи: на договор", "млн ₽", MLN),
            ("admin_fixed", "Административные: постоянная часть", "млн ₽ в год", MLN),
            ("admin_head", "Административные: на сотрудника", "млн ₽ в год", MLN2),
            ("hardware", "Оборудование на нового сотрудника", "млн ₽", MLN2),
            ("startup_costs", "Регистрация, аккредитация, реестр ПО", "млн ₽", MLN),
            ("cost_growth", "Рост прочих затрат", "доля в год", PCT),
            ("contingency", "Резерв", "доля затрат", PCT),
            ("tax", "Налог на прибыль", "доля", PCT),
            ("loss_cap", "Перенос убытка: доля прибыли", "доля", PCT),
            ("ar_days", "Срок оплаты банками", "дней", INT),
            ("discount", "Ставка дисконтирования", "доля", PCT),
            ("tv_growth", "Рост после 2031 года", "доля в год", PCT),
            ("buffer", "Запас к финансированию", "доля", PCT),
        ]
        for key, label, unit, fmt in consts:
            self.const(key, label, unit, f"={src(key)}", fmt)
        self.const("avg_t1", "Клиентов у банка уровня 1, кроме ВТБ", "млн", f"={avg('avg_t1')}", MLN3, "лист «Рынок»")
        self.const("avg_t2", "Клиентов у банка уровня 2", "млн", f"={avg('avg_t2')}", MLN3, "лист «Рынок»")
        for role in roles:
            self.const(
                f"sal_{role['id']}",
                f"Оклад: {role['name']}",
                "тыс. ₽ в месяц",
                f"={salary(role['id'])}",
                INT,
                "лист «План»",
            )
        for deal in self.data["plan"]["deals"]:
            self.series(
                deal["id"],
                deal["name"],
                deal["unit"],
                lambda col, k=deal["id"]: f"={plan(k, col)}",
                NUM1,
                "лист «План»",
                start=f"={plan(deal['id'], START_COL)}",
                total=True,
                kind="link",
            )
        for role in roles:
            self.series(
                f"hc_{role['id']}",
                f"Численность: {role['name']}",
                "ставок",
                lambda col, k=role["id"]: f"={plan(k, col)}",
                NUM1,
                "лист «План»",
                start=f"={plan(role['id'], START_COL)}",
                kind="link",
            )

        self.section("1. Время")
        self.series(
            "year",
            "Год",
            "",
            lambda col: f"={prev(col)}{self.rows['year']}+1",
            YEAR,
            "предыдущий + 1",
            start=f"={start_year}",
        )
        self.series(
            "idx",
            "Индекс цен к 2027 году",
            "раз",
            lambda col: f"=(1+{c('indexation')})^({y(col)}-$E${self.rows['year']}-1)",
            "0.000",
            "(1 + индексация) ^ (год − 2027)",
        )
        self.series(
            "idx_cost",
            "Индекс прочих затрат к 2027 году",
            "раз",
            lambda col: f"=(1+{c('cost_growth')})^({y(col)}-$E${self.rows['year']}-1)",
            "0.000",
            "то же для затрат",
        )
        self.series(
            "idx_sal",
            "Индекс окладов к 2027 году",
            "раз",
            lambda col: f"=(1+{c('raise')})^({y(col)}-$E${self.rows['year']}-1)",
            "0.000",
            "то же для окладов",
        )
        self.series(
            "base",
            "Предельная база взносов",
            "млн ₽ в год",
            lambda col: f"={c('contrib_base')}*(1+{c('base_growth')})^({y(col)}-$E${self.rows['year']})",
            MLN3,
            "база 2026 × (1 + рост) ^ (год − 2026)",
        )

        self.section("2. Договоры")
        self.series(
            "vtb_new",
            "Договор с ВТБ (1 — в этом году)",
            "банков",
            lambda col: f"=IF({y(col)}={c('vtb_year')},1,0)",
            INT,
            "год договора с ВТБ",
        )
        self.series(
            "t1_new",
            "Новые банки уровня 1, с ВТБ",
            "банков",
            lambda col: f"={s('vtb_new', col)}+{s('new_t1', col)}",
            INT,
            "ВТБ + другие",
            start="=0",
            total=True,
        )
        self.series(
            "t1_cum",
            "Банки уровня 1 с договором, на конец года",
            "банков",
            lambda col: f"={prev(col)}{self.rows['t1_cum']}+{s('t1_new', col)}",
            INT,
            "накопленным итогом",
            start="=0",
        )
        self.series(
            "t2_cum",
            "Банки уровня 2 с договором, на конец года",
            "банков",
            lambda col: f"={prev(col)}{self.rows['t2_cum']}+{s('new_t2', col)}",
            INT,
            "накопленным итогом",
            start="=0",
        )
        self.series(
            "banks",
            "Банков с договором, на конец года",
            "банков",
            lambda col: f"={s('t1_cum', col)}+{s('t2_cum', col)}",
            INT,
            "уровень 1 + уровень 2",
            start="=0",
        )
        self.series(
            "deals",
            "Новых договоров за год",
            "банков",
            lambda col: f"={s('t1_new', col)}+{s('new_t2', col)}",
            INT,
            "",
            total=True,
        )

        self.section("3. Выручка, млн ₽ без НДС")
        vtb_sub = f"({c('fixed_t1')}+{c('clients')}*{c('price_client')})"
        t1_sub = f"({c('fixed_t1')}+{c('avg_t1')}*{c('price_client')})"
        t2_sub = f"({c('fixed_t2')}+{c('avg_t2')}*{c('price_client')})"
        yr = self.rows["year"]

        def cohort(sub: str, key: str, col: str, through: bool = False) -> str:
            last = col if through else prev(col)
            shift = "" if through else "-1"
            return (
                f"{sub}*{s('idx', col)}*SUMPRODUCT($E${self.rows[key]}:{last}${self.rows[key]},"
                f"(1+{c('expansion')})^({y(col)}-$E${yr}:{last}${yr}{shift}))"
            )

        self.series(
            "rev_vtb",
            "Подписка ВТБ",
            "млн ₽",
            lambda col: (
                f"=IF({y(col)}>{c('vtb_year')},{vtb_sub}*(1+{c('expansion')})"
                f"^({y(col)}-{c('vtb_year')}-1)*{s('idx', col)},0)"
            ),
            MLN,
            "со следующего года после договора; растёт с новыми операциями; индексируется",
            total=True,
        )
        self.series(
            "rev_t1",
            "Подписка других банков уровня 1",
            "млн ₽",
            lambda col: f"={cohort(t1_sub, 'new_t1', col)}",
            MLN,
            "по годам договоров: подписка × (1 + рост) ^ (лет после первого года) × индекс",
            total=True,
        )
        self.series(
            "rev_t2",
            "Подписка банков уровня 2",
            "млн ₽",
            lambda col: f"={cohort(t2_sub, 'new_t2', col)}",
            MLN,
            "то же для уровня 2",
            total=True,
        )
        self.series(
            "rev_impl",
            "Внедрение",
            "млн ₽",
            lambda col: f"=({s('t1_new', col)}*{c('impl_t1')}+{s('new_t2', col)}*{c('impl_t2')})*{s('idx', col)}",
            MLN,
            "новые договоры × плата за внедрение × индекс",
            total=True,
        )
        self.series(
            "rev_sub",
            "Подписка всего",
            "млн ₽",
            lambda col: f"={s('rev_vtb', col)}+{s('rev_t1', col)}+{s('rev_t2', col)}",
            MLN,
            "",
            total=True,
        )
        self.series(
            "revenue",
            "Выручка",
            "млн ₽",
            lambda col: f"={s('rev_sub', col)}+{s('rev_impl', col)}",
            MLN,
            "подписка + внедрение",
            total=True,
        )
        self.series(
            "arr",
            "Годовая подписка на конец года (ARR)",
            "млн ₽ в год",
            lambda col: (
                f"=IF({y(col)}>={c('vtb_year')},{vtb_sub}*(1+{c('expansion')})^({y(col)}-{c('vtb_year')})"
                f"*{s('idx', col)},0)+{cohort(t1_sub, 'new_t1', col, True)}+{cohort(t2_sub, 'new_t2', col, True)}"
            ),
            MLN,
            "подписка всех банков с договором в ценах года",
        )

        self.section("4. Затраты, млн ₽")
        hc_sum = "+".join(f"{{col}}{self.rows['hc_' + r_['id']]}" for r_ in roles)
        self.series(
            "hc",
            "Численность команды",
            "ставок",
            lambda col: "=" + hc_sum.format(col=col),
            NUM1,
            "сумма по ролям",
            start="=" + hc_sum.format(col="E"),
        )
        self.series(
            "hires",
            "Новые сотрудники",
            "ставок",
            lambda col: f"=MAX(0,{s('hc', col)}-{prev(col)}{self.rows['hc']})",
            NUM1,
            "прирост численности",
        )
        for role in roles:
            rid = role["id"]
            self.series(
                f"pp_{rid}",
                f"Годовой оклад на ставку: {role['name']}",
                "млн ₽",
                lambda col, k=rid: f"={c('sal_' + k)}*12/1000*{s('idx_sal', col)}",
                MLN2,
                "оклад × 12 × индекс",
            )
        for role in roles:
            rid = role["id"]
            self.series(
                f"pay_{rid}",
                f"Фонд оплаты труда: {role['name']}",
                "млн ₽",
                lambda col, k=rid: f"={s('hc_' + k, col)}*{s('pp_' + k, col)}",
                MLN,
                "ставки × годовой оклад",
                total=True,
            )
        self.series(
            "payroll",
            "Фонд оплаты труда",
            "млн ₽",
            lambda col: "=" + "+".join(s(f"pay_{r_['id']}", col) for r_ in roles),
            MLN,
            "",
            total=True,
        )

        def contrib(col: str) -> str:
            parts = [
                f"{s('hc_' + r_['id'], col)}*(MIN({s('pp_' + r_['id'], col)},{s('base', col)})*{c('contrib_low')}"
                f"+MAX(0,{s('pp_' + r_['id'], col)}-{s('base', col)})*{c('contrib_high')}"
                f"+{s('pp_' + r_['id'], col)}*{c('contrib_injury')})"
                for r_ in roles
            ]
            return "=" + "+".join(parts)

        self.series(
            "contrib",
            "Страховые взносы",
            "млн ₽",
            contrib,
            MLN,
            "по каждой роли: в пределах базы × льготный тариф + сверх базы × 7,6 % + травматизм",
            total=True,
        )
        self.series(
            "first",
            "Первый договор (1 — в этом году)",
            "",
            lambda col: f"=IF(AND({s('banks', col)}>0,{prev(col)}{self.rows['banks']}=0),1,0)",
            INT,
            "",
        )
        self.series(
            "cost_sec",
            "Оценка ИБ (анализ уязвимостей)",
            "млн ₽",
            lambda col: (
                f"=({s('first', col)}*{c('sec_first')}+IF({prev(col)}{self.rows['banks']}>0,1,0)*{c('sec_annual')})"
                f"*{s('idx_cost', col)}"
            ),
            MLN,
            "первая — в год первого договора, дальше — ежегодно",
            total=True,
        )
        self.series(
            "cost_legal",
            "Внешняя юридическая экспертиза правил",
            "млн ₽",
            lambda col: f"={c('legal_ext')}*{s('idx_cost', col)}",
            MLN,
            "",
            total=True,
        )
        self.series(
            "cost_infra",
            "Инфраструктура и инструменты",
            "млн ₽",
            lambda col: f"={c('infra_head')}*{s('hc', col)}*{s('idx_cost', col)}",
            MLN,
            "на сотрудника",
            total=True,
        )
        self.series(
            "cost_sales",
            "Продажи и маркетинг",
            "млн ₽",
            lambda col: f"=({c('sales_fixed')}+{c('sales_deal')}*{s('deals', col)})*{s('idx_cost', col)}",
            MLN,
            "постоянная часть + на договор",
            total=True,
        )
        self.series(
            "cost_admin",
            "Административные",
            "млн ₽",
            lambda col: f"=({c('admin_fixed')}+{c('admin_head')}*{s('hc', col)})*{s('idx_cost', col)}",
            MLN,
            "постоянная часть + на сотрудника",
            total=True,
        )
        self.series(
            "cost_hw",
            "Оборудование",
            "млн ₽",
            lambda col: f"={c('hardware')}*{s('hires', col)}*{s('idx_cost', col)}",
            MLN,
            "на нового сотрудника",
            total=True,
        )
        self.series(
            "cost_start",
            "Регистрация, аккредитация, реестр ПО",
            "млн ₽",
            lambda col: f"=IF({y(col)}=$F${yr},{c('startup_costs')},0)",
            MLN,
            "только 2027 год",
            total=True,
        )
        self.series(
            "subtotal",
            "Затраты до резерва",
            "млн ₽",
            lambda col: (
                f"={s('payroll', col)}+{s('contrib', col)}+{s('cost_sec', col)}+{s('cost_legal', col)}"
                f"+{s('cost_infra', col)}+{s('cost_sales', col)}+{s('cost_admin', col)}+{s('cost_hw', col)}"
                f"+{s('cost_start', col)}"
            ),
            MLN,
            "",
            total=True,
        )
        self.series(
            "cost_res",
            "Резерв на непредвиденное",
            "млн ₽",
            lambda col: f"={s('subtotal', col)}*{c('contingency')}",
            MLN,
            "",
            total=True,
        )
        self.series(
            "opex",
            "Затраты всего",
            "млн ₽",
            lambda col: f"={s('subtotal', col)}+{s('cost_res', col)}",
            MLN,
            "",
            total=True,
        )

        self.section("5. Отчёт о прибылях, млн ₽")
        self.series(
            "ebitda",
            "Операционная прибыль (EBITDA)",
            "млн ₽",
            lambda col: f"={s('revenue', col)}-{s('opex', col)}",
            MLN,
            "выручка − затраты; амортизации нет: оборудование списывается в год покупки",
            total=True,
        )
        self.series(
            "margin",
            "Операционная рентабельность",
            "доля",
            lambda col: f"=IF({s('revenue', col)}>0,{s('ebitda', col)}/{s('revenue', col)},0)",
            PCT,
            "",
        )
        self.series(
            "loss_open",
            "Перенесённый убыток на начало года",
            "млн ₽",
            lambda col: f"={prev(col)}{self.rows['loss_open'] + 4}",
            MLN,
            "",
        )
        self.series(
            "loss_used",
            "Использовано убытка",
            "млн ₽",
            lambda col: f"=MIN({s('loss_open', col)},MAX(0,{s('ebitda', col)})*{c('loss_cap')})",
            MLN,
            "не больше доли прибыли года",
        )
        self.series(
            "taxable",
            "Налоговая база",
            "млн ₽",
            lambda col: f"=MAX(0,{s('ebitda', col)})-{s('loss_used', col)}",
            MLN,
            "",
        )
        self.series(
            "tax_paid", "Налог на прибыль", "млн ₽", lambda col: f"={s('taxable', col)}*{c('tax')}", MLN, "", total=True
        )
        self.series(
            "loss_close",
            "Перенесённый убыток на конец года",
            "млн ₽",
            lambda col: f"={s('loss_open', col)}-{s('loss_used', col)}+MAX(0,-{s('ebitda', col)})",
            MLN,
            "",
            start="=0",
        )
        self.series(
            "net",
            "Чистая прибыль",
            "млн ₽",
            lambda col: f"={s('ebitda', col)}-{s('tax_paid', col)}",
            MLN,
            "",
            total=True,
        )

        self.section("6. Деньги, млн ₽")
        self.series(
            "ar",
            "Дебиторская задолженность на конец года",
            "млн ₽",
            lambda col: f"={s('revenue', col)}*{c('ar_days')}/365",
            MLN,
            "выручка × дни / 365",
            start="=0",
        )
        self.series(
            "d_ar",
            "Прирост дебиторской задолженности",
            "млн ₽",
            lambda col: f"={s('ar', col)}-{prev(col)}{self.rows['ar']}",
            MLN,
            "",
            total=True,
        )
        self.series(
            "fcf",
            "Свободный денежный поток",
            "млн ₽",
            lambda col: f"={s('net', col)}-{s('d_ar', col)}",
            MLN,
            "чистая прибыль − прирост дебиторки; капвложений нет",
            total=True,
        )
        self.series(
            "cum",
            "Накопленный денежный поток",
            "млн ₽",
            lambda col: f"={prev(col)}{self.rows['cum']}+{s('fcf', col)}",
            MLN,
            "до внешнего финансирования",
            start="=0",
        )
        self.series(
            "df",
            "Коэффициент дисконтирования",
            "",
            lambda col: f"=1/(1+{c('discount')})^({y(col)}-$E${yr})",
            "0.000",
            "конец года, к концу 2026",
        )
        self.series(
            "pv", "Дисконтированный поток", "млн ₽", lambda col: f"={s('fcf', col)}*{s('df', col)}", MLN, "", total=True
        )
        self.series(
            "tv",
            "Остаточная стоимость (только 2031)",
            "млн ₽",
            lambda col: (
                f"=IF({y(col)}=$J${yr},MAX(0,{s('fcf', col)})*(1+{c('tv_growth')})"
                f"/({c('discount')}-{c('tv_growth')}),0)"
            ),
            MLN,
            "поток 2031 × (1 + рост) / (ставка − рост); отрицательный поток не продлевается",
        )
        self.series(
            "pv_tv",
            "Дисконтированная остаточная стоимость",
            "млн ₽",
            lambda col: f"={s('tv', col)}*{s('df', col)}",
            MLN,
            "",
            total=True,
        )
        self.series(
            "fcf_tv",
            "Поток с остаточной стоимостью",
            "млн ₽",
            lambda col: f"={s('fcf', col)}+{s('tv', col)}",
            MLN,
            "для IRR",
        )
        self.series(
            "f_ebitda",
            "Операционная прибыль впервые (год)",
            "",
            lambda col: (
                f'=IF(AND({s("ebitda", col)}>0,COUNTIF($F${self.rows["ebitda"]}:{col}${self.rows["ebitda"]},">0")=1),'
                f'{y(col)},"")'
            ),
            YEAR,
            "служебная строка",
        )
        self.series(
            "f_cum",
            "Накопленный поток впервые положителен (год)",
            "",
            lambda col: (
                f'=IF(AND({s("cum", col)}>0,COUNTIF($F${self.rows["cum"]}:{col}${self.rows["cum"]},">0")=1),'
                f'{y(col)},"")'
            ),
            YEAR,
            "служебная строка",
        )

        self.section("7. Итоги сценария")
        res = [
            (
                "revenue_total",
                "Выручка 2027–2031",
                "млн ₽",
                f"=SUM(F{self.rows['revenue']}:J{self.rows['revenue']})",
                MLN,
            ),
            ("revenue_2031", "Выручка 2031 года", "млн ₽", f"=J{self.rows['revenue']}", MLN),
            ("arr_2031", "Годовая подписка на конец 2031 года (ARR)", "млн ₽", f"=J{self.rows['arr']}", MLN),
            ("ebitda_2031", "Операционная прибыль 2031 года", "млн ₽", f"=J{self.rows['ebitda']}", MLN),
            ("banks_2031", "Банков с договором на конец 2031 года", "банков", f"=J{self.rows['banks']}", INT),
            ("hc_2031", "Команда на конец 2031 года", "ставок", f"=J{self.rows['hc']}", NUM1),
            (
                "min_cum",
                "Наибольший накопленный дефицит",
                "млн ₽",
                f"=MIN(0,MIN(F{self.rows['cum']}:J{self.rows['cum']}))",
                MLN,
            ),
            ("funding", "Потребность во внешнем финансировании с запасом", "млн ₽", "", MLN),
            ("cum_2031", "Накопленный поток на конец 2031 года", "млн ₽", f"=J{self.rows['cum']}", MLN),
            (
                "npv",
                "NPV 2027–2031 без остаточной стоимости",
                "млн ₽",
                f"=SUM(F{self.rows['pv']}:J{self.rows['pv']})",
                MLN,
            ),
            (
                "npv_tv",
                "NPV с остаточной стоимостью",
                "млн ₽",
                f"=SUM(F{self.rows['pv']}:J{self.rows['pv']})+SUM(F{self.rows['pv_tv']}:J{self.rows['pv_tv']})",
                MLN,
            ),
            (
                "irr",
                "IRR с остаточной стоимостью",
                "доля",
                f'=IFERROR(IRR(F{self.rows["fcf_tv"]}:J{self.rows["fcf_tv"]}),"н/д")',
                PCT,
            ),
            (
                "year_ebitda",
                "Год выхода на операционную прибыль",
                "год",
                f'=IF(COUNT(F{self.rows["f_ebitda"]}:J{self.rows["f_ebitda"]})=0,"после 2031",'
                f"MIN(F{self.rows['f_ebitda']}:J{self.rows['f_ebitda']}))",
                YEAR,
            ),
            (
                "year_cum",
                "Год, когда накопленный поток становится положительным",
                "год",
                f'=IF(COUNT(F{self.rows["f_cum"]}:J{self.rows["f_cum"]})=0,"после 2031",'
                f"MIN(F{self.rows['f_cum']}:J{self.rows['f_cum']}))",
                YEAR,
            ),
        ]
        for key, label, unit, formula, fmt in res:
            if key == "funding":
                formula = f"=-{self.c('min_cum')}*(1+{c('buffer')})"
            self.const(key, label, unit, formula, fmt, kind="calc")
            self.ws[f"D{self.rows[key]}"].font = font(TEXT, True)


def calc_sheets(wb: Workbook, inp: Inputs, plan: Plan, market: Columns, data: dict) -> list[Block]:
    blocks = []
    start = f"Допущения!$D${inp.rows['start_year']}"
    for s, name in enumerate(SHEET_CALC):
        ws = wb.create_sheet(name)
        title(
            ws,
            f"Расчёт: {name.lower()} сценарий",
            "Входы — зелёные ссылки на «Допущения», «План» и «Рынок»; в строке одна формула для всех лет; "
            "млн ₽ без НДС",
        )
        blk = Block(
            ws,
            4,
            data,
            src=lambda key, s=s: inp.ref(key, s),
            plan=lambda key, col, s=s: plan.ref(s, key, col),
            salary=lambda key, s=s: plan.salary(s, key),
            start_year=start,
            avg=lambda key, s=s: market.ext(key, s),
        )
        widths(ws, {"A": 5, "B": 50, "C": 14, "D": 12, "E": 9, **dict.fromkeys(YEAR_COLS, 10), "K": 12, "L": 60})
        ws.freeze_panes = "F5"
        page(ws)
        blocks.append(blk)
    return blocks


# ---------------------------------------------------------------- сводка, чувствительность, проверки


def summary_sheet(wb: Workbook, blocks: list[Block], bank: Columns, market: Columns, data: dict) -> dict[str, int]:
    ws = wb.create_sheet("Сводка", 0)
    title(ws, "Сводка по сценариям", "Все числа — ссылки на листы сценариев; млн ₽ без НДС")
    header(
        ws,
        4,
        {
            "A": "№",
            "B": "Показатель",
            "C": "Единица",
            "D": data["scenarios"][0],
            "E": data["scenarios"][1],
            "F": data["scenarios"][2],
        },
    )
    rows: dict[str, int] = {}
    r = 5
    items = [
        ("Эффект для ВТБ", None),
        ("fte_round", "Высвобождается ставок", "ставок", "bank", INT),
        ("prevented_round", "Предотвращённый ущерб клиентов", "млн ₽ в год", "bank", INT),
        ("effect_capacity", "Прямой эффект банка: ёмкость", "млн ₽ в год", "bank", MLN),
        ("effect_cash", "Прямой эффект банка: деньги", "млн ₽ в год", "bank", MLN),
        ("sub_vtb", "Подписка для ВТБ (цены 2027)", "млн ₽ в год", "bank", MLN),
        ("payback_cash", "Окупаемость для ВТБ: деньги", "месяцев", "bank", MONTHS),
        ("roi_cash", "Отдача за 3 года на вложенный рубль: деньги", "раз", "bank", '0.0"×";(0.0"×");"–"'),
        ("Рынок", None),
        ("tam", "TAM: все банки России", "млн ₽ в год", "market", MLN),
        ("sam", "SAM: банки уровней 1 и 2, одна операция", "млн ₽ в год", "market", MLN),
        ("sam_full", "SAM при полном охвате операций", "млн ₽ в год", "market", MLN),
        ("som", "SOM: подписка на конец 2031 (цены 2027)", "млн ₽ в год", "market", MLN),
        ("som_share", "SOM в доле SAM при полном охвате", "доля", "market", PCT),
        ("Проект", None),
        ("revenue_total", "Выручка 2027–2031", "млн ₽", "calc", MLN),
        ("revenue_2031", "Выручка 2031", "млн ₽", "calc", MLN),
        ("arr_2031", "ARR на конец 2031", "млн ₽", "calc", MLN),
        ("ebitda_2031", "Операционная прибыль 2031", "млн ₽", "calc", MLN),
        ("banks_2031", "Банков с договором к концу 2031", "банков", "calc", INT),
        ("hc_2031", "Команда к концу 2031", "ставок", "calc", NUM1),
        ("funding", "Потребность в финансировании с запасом", "млн ₽", "calc", MLN),
        ("year_ebitda", "Год выхода на операционную прибыль", "год", "calc", YEAR),
        ("year_cum", "Год окупаемости вложений", "год", "calc", YEAR),
        ("npv", "NPV без остаточной стоимости", "млн ₽", "calc", MLN),
        ("npv_tv", "NPV с остаточной стоимостью", "млн ₽", "calc", MLN),
        ("irr", "IRR с остаточной стоимостью", "доля", "calc", PCT),
    ]
    for item in items:
        if item[1] is None:
            section(ws, r, item[0], "F")
            r += 1
            continue
        key, label, unit, where, fmt = item
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = len(rows) + 1, label, unit
        ws[f"A{r}"].font, ws[f"B{r}"].font, ws[f"C{r}"].font = font(MUTED), font(TEXT), font(MUTED)
        for s, col in enumerate(SCEN_COLS):
            if where == "bank":
                ref = bank.ext(key, s)
            elif where == "market":
                ref = market.ext(key, s)
            else:
                ref = f"{SHEET_CALC[s]}!$D${blocks[s].rows[key]}"
            ws[f"{col}{r}"] = f"={ref}"
            style_value(ws[f"{col}{r}"], "link", fmt)
        for col in "ABCDEF":
            ws[f"{col}{r}"].border = BOTTOM
        rows[key] = r
        r += 1

    r += 1
    series = [
        ("revenue", "Выручка"),
        ("ebitda", "Операционная прибыль"),
        ("cum", "Накопленный денежный поток"),
        ("banks", "Банков с договором"),
        ("hc", "Команда, ставок"),
    ]
    header(
        ws,
        r,
        {
            "A": "№",
            "B": "Показатель по годам",
            "C": "Сценарий",
            **{c: str(y) for c, y in zip(["D", "E", "F", "G", "H"], data["years"], strict=True)},
        },
    )
    r += 1
    for key, label in series:
        for s, scen in enumerate(data["scenarios"]):
            ws[f"B{r}"], ws[f"C{r}"] = label, scen
            ws[f"B{r}"].font, ws[f"C{r}"].font = font(TEXT), font(MUTED)
            for i, col in enumerate(["D", "E", "F", "G", "H"]):
                ws[f"{col}{r}"] = f"={SHEET_CALC[s]}!{YEAR_COLS[i]}{blocks[s].rows[key]}"
                style_value(ws[f"{col}{r}"], "link", INT if key == "banks" else (NUM1 if key == "hc" else MLN))
            for col in "ABCDEFGH":
                ws[f"{col}{r}"].border = BOTTOM
            rows[f"{key}_{s}"] = r
            r += 1

    def chart(kind: type, key: str, label: str, anchor: str) -> None:
        ch = kind()
        ch.title = label
        ch.y_axis.title = "млн ₽"
        ch.height, ch.width = 7.5, 16
        for s, scen in enumerate(data["scenarios"]):
            row = rows[f"{key}_{s}"]
            ch.series.append(Series(Reference(ws, min_col=4, max_col=8, min_row=row, max_row=row), title=scen))
        header_row = rows["revenue_0"] - 1
        ch.set_categories(Reference(ws, min_col=4, max_col=8, min_row=header_row, max_row=header_row))
        ws.add_chart(ch, anchor)

    chart(BarChart, "revenue", "Выручка по сценариям, млн ₽", "J4")
    chart(LineChart, "cum", "Накопленный денежный поток, млн ₽", "J20")
    widths(ws, {"A": 5, "B": 46, "C": 16, "D": 15, "E": 15, "F": 15, "G": 12, "H": 12})
    ws.freeze_panes = "D5"
    page(ws)
    return rows


def sensitivity_sheet(wb: Workbook, inp: Inputs, plan: Plan, market: Columns, data: dict, base: Block) -> dict:
    ws = wb.create_sheet("Чувствительность")
    sens = data["sensitivity"]
    step = sens["step"]
    title(
        ws,
        "Чувствительность NPV базового сценария",
        f"Каждое допущение меняется на ±{int(step * 100)} %, остальные — как в базовом сценарии. "
        "Блоки расчёта ниже — та же "
        "структура, что на листе «Базовый»",
    )
    header(
        ws,
        4,
        {
            "A": "№",
            "B": "Допущение",
            "C": "Изменение",
            "D": f"NPV при −{int(step * 100)} %",
            "E": f"NPV при +{int(step * 100)} %",
            "F": "Базовый NPV",
            "G": "Размах",
        },
    )
    start = f"Допущения!$D${inp.rows['start_year']}"
    drivers = sens["drivers"]
    table_rows = {d["id"]: 5 + i for i, d in enumerate(drivers)}
    shift_row = 5 + len(drivers)
    top = shift_row + 4
    blocks: dict[str, Block] = {}
    s = 1

    def make(
        tag: str,
        label: str,
        inputs: set[str],
        plan_keys: set[str],
        salaries: bool,
        roles: bool,
        mult: float,
        add: dict[str, float] | None = None,
    ) -> Block:
        nonlocal top
        ws[f"A{top}"] = label
        ws[f"A{top}"].font = font(NAVY, True)
        add = add or {}

        def src(key: str) -> str:
            ref = inp.ref(key, s)
            if key in inputs:
                return f"{ref}*{mult}"
            if key in add:
                return f"{ref}+{add[key]}"
            return ref

        def pl(key: str, col: str) -> str:
            ref = plan.ref(s, key, col)
            is_role = any(r_["id"] == key for r_ in data["plan"]["roles"])
            if key in plan_keys or (roles and is_role):
                return f"{ref}*{mult}"
            return ref

        def sal(key: str) -> str:
            return f"{plan.salary(s, key)}*{mult}" if salaries else plan.salary(s, key)

        blk = Block(
            ws, top + 1, data, src=src, plan=pl, salary=sal, start_year=start, avg=lambda key: market.ext(key, s)
        )
        for row in range(top, blk.r):
            ws.row_dimensions[row].outlineLevel = 1
            ws.row_dimensions[row].hidden = True
        top = blk.r + 2
        blocks[tag] = blk
        return blk

    ref_blk = make("base", "Контроль: базовый сценарий без изменений", set(), set(), False, False, 1)
    for i, d in enumerate(drivers):
        r = table_rows[d["id"]]
        ins = set(d.get("inputs", []))
        pks = set(d.get("plan", []))
        lo = make(
            f"{d['id']}_lo",
            f"{d['name']}: −{int(step * 100)} %",
            ins,
            pks,
            bool(d.get("salaries")),
            bool(d.get("roles")),
            round(1 - step, 6),
        )
        hi = make(
            f"{d['id']}_hi",
            f"{d['name']}: +{int(step * 100)} %",
            ins,
            pks,
            bool(d.get("salaries")),
            bool(d.get("roles")),
            round(1 + step, 6),
        )
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = i + 1, d["name"], f"±{int(step * 100)} %"
        ws[f"D{r}"] = f"=$D${lo.rows['npv_tv']}"
        ws[f"E{r}"] = f"=$D${hi.rows['npv_tv']}"
        ws[f"F{r}"] = f"=Базовый!$D${base.rows['npv_tv']}"
        ws[f"G{r}"] = f"=ABS(E{r}-D{r})"
        for col in "DEFG":
            style_value(ws[f"{col}{r}"], "calc" if col in "DEG" else "link", MLN)
        for col in "ABCDEFG":
            ws[f"{col}{r}"].border = BOTTOM
        ws[f"B{r}"].font = font(TEXT)
    sh = sens["shift"]
    r = shift_row
    blk = make(sh["id"], sh["name"], set(), set(), False, False, 1, add={sh["input"]: sh["add"]})
    ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = len(drivers) + 1, sh["name"], f"+{sh['add']} год"
    ws[f"D{r}"] = f"=$D${blk.rows['npv_tv']}"
    ws[f"F{r}"] = f"=Базовый!$D${base.rows['npv_tv']}"
    ws[f"G{r}"] = f"=ABS(D{r}-F{r})"
    for col in "DFG":
        style_value(ws[f"{col}{r}"], "calc" if col != "F" else "link", MLN)
    for col in "ABCDEFG":
        ws[f"{col}{r}"].border = BOTTOM
    ws[f"A{shift_row + 2}"] = "Строки с блоками расчёта свёрнуты: раскройте группу слева, чтобы увидеть формулы."
    ws[f"A{shift_row + 2}"].font = font(MUTED, size=9, italic=True)
    widths(
        ws, {"A": 5, "B": 50, "C": 14, "D": 14, "E": 14, "F": 14, "G": 12, "H": 10, "I": 10, "J": 10, "K": 12, "L": 50}
    )
    ws.sheet_properties.outlinePr.summaryBelow = False
    page(ws)
    return {"rows": table_rows, "shift_row": shift_row, "ref_block": ref_blk, "drivers": drivers, "shift": sh}


def checks_sheet(wb: Workbook, blocks: list[Block], bank: Columns, market: Columns, sens: dict, data: dict) -> dict:
    ws = wb.create_sheet("Проверки")
    title(
        ws,
        "Проверки",
        "Целостность — должна быть «ОК» всегда; допущения — предупреждение, если цена или доля рынка выходят за рамки",
    )
    header(
        ws,
        4,
        {"A": "№", "B": "Проверка", "C": "Сценарий", "D": "Ожидается", "E": "Факт", "F": "Статус", "G": "Пояснение"},
    )
    r = 5
    n = 0
    integrity: list[str] = []
    warnings: list[str] = []
    registry = registry_values()

    def add(
        label: str, scen: str, expected: str, actual: str, cond: str, note: str, kind: str = "integrity", fmt: str = MLN
    ) -> None:
        nonlocal r, n
        n += 1
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"G{r}"] = n, label, scen, note
        ws[f"D{r}"] = expected
        ws[f"E{r}"] = actual
        style_value(ws[f"D{r}"], "input" if not str(expected).startswith("=") else "calc", fmt)
        style_value(ws[f"E{r}"], "calc", fmt)
        bad = "ОШИБКА" if kind == "integrity" else "ВНИМАНИЕ"
        ws[f"F{r}"] = f'=IF({cond},"ОК","{bad}")'
        ws[f"F{r}"].font = font(TEXT, True)
        ws[f"F{r}"].alignment = Alignment(horizontal="center")
        for col in "ABCDEFG":
            ws[f"{col}{r}"].border = BOTTOM
        ws[f"B{r}"].font, ws[f"G{r}"].font = font(TEXT), font(MUTED, size=9)
        (integrity if kind == "integrity" else warnings).append(f"F{r}")
        r += 1

    section(ws, r, "Целостность", "G")
    r += 1
    for s, scen in enumerate(data["scenarios"]):
        add(
            "Ставки совпадают с презентацией (C-15)",
            scen,
            registry["fte_round"][s],
            f"={bank.ext('fte_round', s)}",
            f"E{r}=D{r}",
            "эталон — реестр утверждений",
            fmt=INT,
        )
        add(
            "Ущерб совпадает с презентацией (C-16)",
            scen,
            registry["prevented_round"][s],
            f"={bank.ext('prevented_round', s)}",
            f"E{r}=D{r}",
            "эталон — реестр утверждений",
            fmt=INT,
        )
    for s, scen in enumerate(data["scenarios"]):
        b = blocks[s]
        sh = SHEET_CALC[s]
        add(
            "Выручка = подписка + внедрение",
            scen,
            f"={sh}!K{b.rows['rev_sub']}+{sh}!K{b.rows['rev_impl']}",
            f"={sh}!K{b.rows['revenue']}",
            f"ABS(E{r}-D{r})<0.000001",
            "итоги 2027–2031",
        )
        add(
            "Подписка = ВТБ + уровень 1 + уровень 2",
            scen,
            f"={sh}!K{b.rows['rev_vtb']}+{sh}!K{b.rows['rev_t1']}+{sh}!K{b.rows['rev_t2']}",
            f"={sh}!K{b.rows['rev_sub']}",
            f"ABS(E{r}-D{r})<0.000001",
            "",
        )
        add(
            "Накопленный поток = сумма потоков",
            scen,
            f"={sh}!K{b.rows['fcf']}",
            f"={sh}!J{b.rows['cum']}",
            f"ABS(E{r}-D{r})<0.000001",
            "",
        )
        add(
            "Чистая прибыль = EBITDA − налог",
            scen,
            f"={sh}!K{b.rows['ebitda']}-{sh}!K{b.rows['tax_paid']}",
            f"={sh}!K{b.rows['net']}",
            f"ABS(E{r}-D{r})<0.000001",
            "",
        )
        add(
            "Численность и договоры не отрицательные",
            scen,
            0,
            f"=MIN({sh}!E{b.rows['hc']}:J{b.rows['hc']},{sh}!F{b.rows['banks']}:J{b.rows['banks']})",
            f"E{r}>=D{r}",
            "",
            fmt=NUM1,
        )
        add(
            "Банков с договором не больше, чем в SAM",
            scen,
            f"={market.ext('banks_t1', s)}+{market.ext('banks_t2', s)}",
            f"={sh}!J{b.rows['banks']}",
            f"E{r}<=D{r}",
            "",
            fmt=INT,
        )
    for s, scen in enumerate(data["scenarios"]):
        add(
            "Доли уровней в сумме 100 %",
            scen,
            1,
            f"={market.ext('t1_rel_share', s)}+{market.ext('t2_rel_share', s)}"
            f"+{market.ext('rel_t3', s)}/{market.ext('rel_total', s)}",
            f"ABS(E{r}-D{r})<0.000001",
            "",
            fmt=PCT,
        )
        add("Банков уровня 3 не меньше нуля", scen, 0, f"={market.ext('banks_t3', s)}", f"E{r}>=D{r}", "", fmt=INT)
    add(
        "Блок чувствительности без изменений = лист «Базовый»",
        data["scenarios"][1],
        f"=Базовый!$D${blocks[1].rows['npv_tv']}",
        f"=Чувствительность!$D${sens['ref_block'].rows['npv_tv']}",
        f"ABS(E{r}-D{r})<0.000001",
        "блоки чувствительности считаются той же формулой",
    )

    section(ws, r, "Допущения о цене и рынке", "G")
    r += 1
    for s, scen in enumerate(data["scenarios"]):
        add(
            "Подписка не больше 15 % прямого эффекта банка (ёмкость)",
            scen,
            0.15,
            f"={bank.ext('price_share', s)}",
            f"E{r}<=D{r}",
            "клиент оставляет себе не меньше 85 % эффекта",
            kind="warn",
            fmt=PCT,
        )
        add(
            "Купить дешевле, чем сделать самим",
            scen,
            1,
            f"={bank.ext('buy_vs_build', s)}",
            f"E{r}<D{r}",
            "затраты в год за горизонт сравнения",
            kind="warn",
            fmt=PCT,
        )
        add(
            "SOM не больше 40 % SAM при полном охвате",
            scen,
            0.4,
            f"={market.ext('som_share', s)}",
            f"E{r}<=D{r}",
            "доля сегмента за 5 лет",
            kind="warn",
            fmt=PCT,
        )

    r += 1
    ws[f"B{r}"] = "Ошибок целостности"
    ws[f"E{r}"] = "=" + "+".join(f'IF({c}="ОК",0,1)' for c in integrity)
    ws[f"B{r + 1}"] = "Предупреждений о допущениях"
    ws[f"E{r + 1}"] = "=" + "+".join(f'IF({c}="ОК",0,1)' for c in warnings)
    ws[f"B{r + 2}"] = "Итог"
    ws[f"F{r + 2}"] = f'=IF(E{r}=0,"ОК","ОШИБКА")'
    for rr in (r, r + 1, r + 2):
        ws[f"B{rr}"].font = font(NAVY, True)
        style_value(ws[f"E{rr}"], "calc", INT)
    ws[f"F{r + 2}"].font = font(NAVY, True)
    widths(ws, {"A": 5, "B": 56, "C": 16, "D": 12, "E": 12, "F": 12, "G": 50})
    ws.conditional_formatting.add(f"F5:F{r + 2}", _rule("ОК", OK_FILL))
    ws.conditional_formatting.add(f"F5:F{r + 2}", _rule("ОШИБКА", BAD_FILL))
    ws.conditional_formatting.add(f"F5:F{r + 2}", _rule("ВНИМАНИЕ", WARN_FILL))
    ws.freeze_panes = "B5"
    page(ws)
    return {
        "errors": f"Проверки!$E${r}",
        "warnings": f"Проверки!$E${r + 1}",
        "status": f"Проверки!$F${r + 2}",
        "integrity": len(integrity),
        "warn": len(warnings),
    }


def _rule(text: str, color: str):
    from openpyxl.formatting.rule import CellIsRule

    return CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill(color))


def sources_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Источники")
    title(ws, "Источники", "Строки M-xx и C-xx — в CLAIMS_REGISTRY.md репозитория; там же граница каждого утверждения")
    header(ws, 4, {"A": "Код", "B": "Что подтверждает", "C": "Источник", "D": "Дата", "E": "Статус", "F": "Ссылка"})
    for i, src in enumerate(data["sources"], start=5):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"E{i}"], ws[f"F{i}"] = (
            src["id"],
            src["what"],
            src["source"],
            src["date"],
            src["status"],
            src["url"],
        )
        for col in "ABCDEF":
            ws[f"{col}{i}"].font = font(TEXT if col in "AB" else MUTED, size=9 if col not in "AB" else 10)
            ws[f"{col}{i}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws[f"{col}{i}"].border = BOTTOM
        ws[f"F{i}"].hyperlink = src["url"]
        ws[f"F{i}"].font = font(COBALT, size=9)
    widths(ws, {"A": 6, "B": 60, "C": 60, "D": 12, "E": 13, "F": 50})
    ws.freeze_panes = "A5"
    page(ws)


def cover_sheet(wb: Workbook, data: dict, checks: dict, summary: dict[str, int]) -> None:
    ws = wb.create_sheet("Обложка", 0)
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Мандат — финансовая модель"
    ws["B2"].font = font(NAVY, True, 18)
    ws["B3"] = (
        f"Версия {data['version']} от {data['as_of']}. Трек ВТБ «Финтех без границ», «Академия Финтеха 2026». "
        "Деньги — млн ₽ без НДС; цены 2027 года, дальше индексация"
    )
    ws["B3"].font = font(MUTED, size=10)

    r = 5
    ws[f"B{r}"] = "Главное"
    ws[f"B{r}"].font = font(NAVY, True, 12)
    r += 1
    header(ws, r, {"B": "Показатель", "C": data["scenarios"][0], "D": data["scenarios"][1], "E": data["scenarios"][2]})
    for key, label, fmt in [
        ("fte_round", "Высвобождается ставок в ВТБ", INT),
        ("prevented_round", "Предотвращённый ущерб клиентов ВТБ, млн ₽ в год", INT),
        ("payback_cash", "Окупаемость для ВТБ (деньги), месяцев", MONTHS),
        ("sam", "SAM: банки уровней 1 и 2, одна операция, млн ₽ в год", MLN),
        ("revenue_2031", "Выручка 2031 года, млн ₽", MLN),
        ("funding", "Потребность в финансировании, млн ₽", MLN),
        ("year_ebitda", "Год выхода на операционную прибыль", YEAR),
        ("npv_tv", "NPV с остаточной стоимостью, млн ₽", MLN),
    ]:
        r += 1
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = font(TEXT)
        ws[f"B{r}"].border = BOTTOM
        for s, col in enumerate(["C", "D", "E"]):
            ws[f"{col}{r}"] = f"=Сводка!{SCEN_COLS[s]}{summary[key]}"
            style_value(ws[f"{col}{r}"], "link", fmt)
            ws[f"{col}{r}"].border = BOTTOM
    r += 2
    ws[f"B{r}"] = "Статус проверок"
    ws[f"B{r}"].font = font(NAVY, True, 12)
    ws[f"C{r}"] = f"={checks['status']}"
    ws[f"C{r}"].font = font(NAVY, True, 12)
    ws[f"D{r}"] = f'="Ошибок: "&{checks["errors"]}&", предупреждений: "&{checks["warnings"]}'
    ws[f"D{r}"].font = font(MUTED)

    r += 2
    ws[f"B{r}"] = "Что внутри"
    ws[f"B{r}"].font = font(NAVY, True, 12)
    sheets = [
        ("Сводка", "главные результаты трёх сценариев и графики"),
        ("Допущения", "все параметры по сценариям с типом, источником и способом проверки на пилоте"),
        ("План", "договоры с банками и команда по годам"),
        ("Эффект для банка", "ставки и ущерб как в презентации, деньги для ВТБ, окупаемость, «купить или сделать»"),
        ("Рынок", "TAM, SAM, SOM снизу вверх и сверху вниз"),
        ("Консервативный, Базовый, Оптимистичный", "полный расчёт сценария: выручка, затраты, отчёт, деньги, NPV"),
        ("Чувствительность", "NPV базового сценария при ±20 % по каждому допущению"),
        ("Проверки", "целостность модели и рамки допущений"),
        ("Источники", "откуда взяты факты"),
    ]
    for name, what in sheets:
        r += 1
        ws[f"B{r}"], ws[f"C{r}"] = name, what
        ws[f"B{r}"].font, ws[f"C{r}"].font = font(TEXT, True), font(TEXT)

    r += 2
    ws[f"B{r}"] = "Как пользоваться"
    ws[f"B{r}"].font = font(NAVY, True, 12)
    howto = [
        "Меняйте только синие ячейки на листах «Допущения» и «План»: каждый сценарий — свой столбец. "
        "Всё остальное — формулы.",
        "Жёлтая заливка — допущения, которые банк заменит своими данными на пилоте "
        "(клиенты, события, минуты, оклады, цена).",
        "Зелёный — ссылка на другой лист, чёрный — формула на этом листе.",
        "После правки посмотрите лист «Проверки»: итог должен быть «ОК».",
        "Пересобрать с нуля: python submission/analytics/build_model.py (нужен LibreOffice для пересчёта).",
    ]
    for text in howto:
        r += 1
        ws[f"B{r}"] = f"• {text}"
        ws[f"B{r}"].font = font(TEXT)

    r += 2
    ws[f"B{r}"] = "Легенда"
    ws[f"B{r}"].font = font(NAVY, True, 12)
    legend = [
        ("12,3", "введено руками (допущение)", INPUT, None),
        ("12,3", "ключевое допущение", INPUT, KEY),
        ("12,3", "ссылка на другой лист", LINK, None),
        ("12,3", "формула", "000000", None),
    ]
    for sample, text, color, bg in legend:
        r += 1
        ws[f"C{r}"], ws[f"D{r}"] = sample, text
        ws[f"C{r}"].font = font(color)
        ws[f"C{r}"].alignment = Alignment(horizontal="center")
        if bg:
            ws[f"C{r}"].fill = fill(bg)
        ws[f"D{r}"].font = font(TEXT)
    widths(ws, {"A": 2, "B": 52, "C": 22, "D": 22, "E": 22})
    page(ws)


# ---------------------------------------------------------------- пересчёт и значения


def recalc(path: Path) -> None:
    """Пересчитывает формулы LibreOffice и записывает значения в тот же файл."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise SystemExit("нужен LibreOffice (soffice) для пересчёта формул")
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        cmd = [
            soffice,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            tmp,
            str(path),
        ]
        subprocess.run(cmd, check=True, capture_output=True, timeout=180)  # noqa: S603
        shutil.copyfile(Path(tmp) / path.name, path)


def read_values(
    path: Path, blocks: list[Block], bank: Columns, market: Columns, sens: dict, checks: dict, data: dict
) -> dict:
    wb = load_workbook(path, data_only=True)
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#") and cell.value.rstrip("!?0/").isupper():
                    errors.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
    out: dict = {
        "scenarios": data["scenarios"],
        "years": data["years"],
        "version": data["version"],
        "as_of": data["as_of"],
        "errors": errors,
    }
    calc = []
    for s, blk in enumerate(blocks):
        ws = wb[SHEET_CALC[s]]
        series = {
            k: [ws[f"{c}{r}"].value for c in YEAR_COLS]
            for k, r in blk.rows.items()
            if ws[f"F{r}"].value is not None and not str(k).startswith("f_")
        }
        consts = {k: ws[f"D{r}"].value for k, r in blk.rows.items() if ws[f"D{r}"].value is not None}
        calc.append({"series": series, "const": consts})
    out["calc"] = calc
    for name, sheet in (("bank", bank), ("market", market)):
        ws = wb[sheet.name]
        out[name] = {k: [ws[f"{c}{r}"].value for c in SCEN_COLS] for k, r in sheet.rows.items()}
    ws = wb["Чувствительность"]
    tornado = []
    for d in sens["drivers"]:
        r = sens["rows"][d["id"]]
        tornado.append(
            {
                "id": d["id"],
                "name": d["name"],
                "low": ws[f"D{r}"].value,
                "high": ws[f"E{r}"].value,
                "base": ws[f"F{r}"].value,
            }
        )
    r = sens["shift_row"]
    tornado.append(
        {
            "id": sens["shift"]["id"],
            "name": sens["shift"]["name"],
            "low": ws[f"D{r}"].value,
            "high": None,
            "base": ws[f"F{r}"].value,
        }
    )
    out["tornado"] = tornado
    chk = wb["Проверки"]
    out["checks"] = {
        "errors": chk[checks["errors"].split("!")[1].replace("$", "")].value,
        "warnings": chk[checks["warnings"].split("!")[1].replace("$", "")].value,
        "status": chk[checks["status"].split("!")[1].replace("$", "")].value,
        "integrity_total": checks["integrity"],
        "warn_total": checks["warn"],
        "rows": [
            {"label": chk[f"B{r}"].value, "scenario": chk[f"C{r}"].value, "status": chk[f"F{r}"].value}
            for r in range(5, chk.max_row + 1)
            if chk[f"F{r}"].value in ("ОК", "ОШИБКА", "ВНИМАНИЕ") and chk[f"B{r}"].value and chk[f"C{r}"].value
        ],
    }
    out["inputs"] = {}
    for sec in data["sections"]:
        for item in sec["items"]:
            out["inputs"][item["id"]] = {
                "values": item["values"],
                "unit": item["unit"],
                "name": item["name"],
                "type": item["type"],
                "source": item["source"],
            }
    out["plan"] = data["plan"]
    out["sources"] = data["sources"]
    return out


def build(recalculate: bool = True) -> dict | None:
    data, _ = load()
    wb = Workbook()
    wb.remove(wb.active)
    inp = Inputs(wb, data)
    plan = Plan(wb, data)
    bank = bank_sheet(wb, inp, data)
    market = market_sheet(wb, inp, bank, data)
    blocks = calc_sheets(wb, inp, plan, market, data)
    # Ссылки рынка на строки сценариев известны только после построения блоков.
    for s in range(3):
        for key, row_key in (("som_nominal", "arr"), ("index_2031", "idx")):
            cell = market.ws[f"{SCEN_COLS[s]}{market.rows[key]}"]
            cell.value = cell.value.replace(f"{{{row_key}}}", str(blocks[s].rows[row_key]))
    summary = summary_sheet(wb, blocks, bank, market, data)
    sens = sensitivity_sheet(wb, inp, plan, market, data, blocks[1])
    checks = checks_sheet(wb, blocks, bank, market, sens, data)
    sources_sheet(wb, data)
    cover_sheet(wb, data, checks, summary)
    order = [
        "Обложка",
        "Сводка",
        "Допущения",
        "План",
        "Эффект для банка",
        "Рынок",
        *SHEET_CALC,
        "Чувствительность",
        "Проверки",
        "Источники",
    ]
    wb._sheets = [wb[n] for n in order]  # noqa: SLF001
    tabs = {
        "Обложка": NAVY,
        "Сводка": NAVY,
        "Допущения": INPUT,
        "План": INPUT,
        "Проверки": "1A9F68",
        "Источники": MUTED,
    }
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = tabs.get(ws.title, COBALT)
    wb.properties.creator = "Команда «Мандата»"
    wb.properties.title = "Мандат — финансовая модель"
    wb.properties.subject = "Эффект для ВТБ, рынок, экономика проекта"
    wb.properties.keywords = "Мандат; МЧД; полномочия; ВТБ; финансовая модель"
    wb.properties.language = "ru-RU"
    OUT.mkdir(exist_ok=True)
    wb.save(XLSX)
    if not recalculate:
        return None
    recalc(XLSX)
    values = read_values(XLSX, blocks, bank, market, sens, checks, data)
    VALUES.write_text(json.dumps(values, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    return values


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--no-recalc", action="store_true", help="не пересчитывать в LibreOffice")
    args = parser.parse_args()
    values = build(recalculate=not args.no_recalc)
    if values is None:
        print(f"записано без пересчёта: {XLSX}")
        return 0
    if values["errors"]:
        sys.stderr.write("ошибки в формулах:\n" + "\n".join(values["errors"][:50]) + "\n")
        return 1
    chk = values["checks"]
    print(
        f"модель: {XLSX.relative_to(ROOT)}; проверки: {chk['status']}, ошибок {chk['errors']}, "
        f"предупреждений {chk['warnings']}"
    )
    for s, name in enumerate(values["scenarios"]):
        c = values["calc"][s]["const"]
        print(
            f"  {name}: выручка 2031 {c['revenue_2031']:.1f}, финансирование {c['funding']:.1f}, "
            f"операционная прибыль с {c['year_ebitda']}, NPV {c['npv_tv']:.1f}"
        )
    return 0 if chk["status"] == "ОК" else 1


if __name__ == "__main__":
    raise SystemExit(main())
